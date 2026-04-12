from datetime import datetime
from typing import List

from openpyxl import load_workbook


REQUIRED_HEADERS = [
    "Date",
    "Amount",
    "Description",
    "Primary Tag",
    "Secondary Tag",
    "Payment Source",
]


def parse_expense_template(file_path: str) -> List[dict]:
    workbook = load_workbook(filename=file_path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    if any(header not in headers for header in REQUIRED_HEADERS):
        raise ValueError(f"Excel template must include headers: {REQUIRED_HEADERS}")

    header_map = {header: index for index, header in enumerate(headers)}
    parsed = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or all(cell is None for cell in row):
            continue

        date_value = row[header_map["Date"]]
        if isinstance(date_value, datetime):
            date_value = date_value.date()
        elif isinstance(date_value, str):
            date_value = datetime.fromisoformat(date_value).date()

        parsed.append(
            {
                "amount": float(row[header_map["Amount"]] or 0),
                "date": date_value,
                "description": row[header_map["Description"]] or "",
                "primary_tag": row[header_map["Primary Tag"]] or "Misc",
                "secondary_tag": row[header_map["Secondary Tag"]] or "Need",
                "payment_source": row[header_map["Payment Source"]] or "Cash",
            }
        )

    return parsed
